"""
Master Schedule Builder — Don Bosco Prep 2026-27
Engine v4: Built entirely from DATA_STRUCTURE.md

Priority system: Raw (fixed) + Total (recalculates per run)
Five entities scored: Student, Teacher, Room, Course Section, Co-Schedule Group
Governing principle: Most restricted gets placed first. Highest score wins conflicts.
"""

import json
import sys
import os
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# POINT VALUES (from DATA_STRUCTURE.md — Actual Point Values)
# ---------------------------------------------------------------------------

GRADE_POINTS = {12: 40, 11: 30, 10: 20, 9: 10}
COHORT_POINTS = 50
SSP_POINTS = 25  # Special Student Population

COURSE_CHAR_POINTS = {
    "ap": 30,
    "singleton": 25,
    "grad_req": 20,
    "semester_only": 15,
    "cohort_course": 15,
    "co_schedule": 15,
    "prescribed_term": 10,
}

LOCK_WEIGHT = 10
ROOM_DEMAND_WEIGHT = 5

PERIODS = ["A", "B", "C", "D", "E", "F", "G"]
TERMS = ["FY", "S1", "S2"]


# ===================================================================
# DATA CLASSES
# ===================================================================

class Student:
    def __init__(self, sid, last_name, first_name, grade, ncaa,
                 leo_ii, leo_i, academic_support, pathway,
                 cohort_name, cohort_locked):
        self.id = sid
        self.last_name = last_name
        self.first_name = first_name
        self.grade = grade
        self.ncaa = ncaa
        self.leo_ii = leo_ii
        self.leo_i = leo_i
        self.academic_support = academic_support
        self.pathway = pathway
        self.cohort_name = cohort_name
        self.cohort_locked = cohort_locked

        self.transcript = []  # list of dicts from Sheet 2
        self.course_requests = []  # list of course codes
        self.raw_priority = 0
        self.total_priority = 0
        self.course_priority_values = {}  # {course_code: points}
        self.placed_courses = set()
        self.priority_history = {}  # {run_number: total_priority}

    def calculate_raw(self):
        val = GRADE_POINTS.get(self.grade, 0)
        if self.leo_ii:
            val += COHORT_POINTS
        ssp = self.leo_i or self.academic_support or self.pathway
        if ssp:
            if self.leo_ii and self.leo_i:
                pass  # LEO I -> LEO II same program, no stacking
            else:
                val += SSP_POINTS
        self.raw_priority = val

    def calculate_total(self):
        remaining = [cr for cr in self.course_requests if cr not in self.placed_courses]
        course_points = sum(self.course_priority_values.get(cr, 0) for cr in remaining)
        self.total_priority = self.raw_priority + course_points

    @property
    def name(self):
        return f"{self.last_name}, {self.first_name}"


class Teacher:
    def __init__(self, tid, last_name, first_name, department,
                 max_periods, sixth_fy, sixth_s1, sixth_s2,
                 avail, ssp_teacher, co_schedule_approved):
        self.id = tid
        self.last_name = last_name
        self.first_name = first_name
        self.department = department
        self.max_periods = max_periods
        self.sixth_fy = sixth_fy
        self.sixth_s1 = sixth_s1
        self.sixth_s2 = sixth_s2
        self.availability = avail  # {period: True/False}
        self.ssp_teacher = ssp_teacher
        self.co_schedule_approved = co_schedule_approved

        self.course_assignments = []  # list of dicts from Sheet 2
        self.raw_priority = 0
        self.total_priority = 0
        self.top_student_value = 0
        self.prescribed_room_value = 0
        self.lock_count = 0
        self.priority_history = {}

    def count_locks(self):
        locks = 0
        for ca in self.course_assignments:
            if ca["prescribed_room"]:
                locks += 1
            if ca["prescribed_period"]:
                locks += 1
            if ca["prescribed_term"]:
                locks += 1
            if ca["prescribed_cohort"]:
                locks += 1
        for period, available in self.availability.items():
            if not available:
                locks += 1
        self.lock_count = locks

    def calculate_raw(self):
        self.count_locks()
        self.raw_priority = self.lock_count * LOCK_WEIGHT

    def calculate_total(self, students, rooms):
        self.top_student_value = 0
        all_course_codes = {ca["course_code"] for ca in self.course_assignments}
        for s in students.values():
            remaining = [cr for cr in s.course_requests if cr not in s.placed_courses]
            if any(cr in all_course_codes for cr in remaining):
                if s.total_priority > self.top_student_value:
                    self.top_student_value = s.total_priority

        self.prescribed_room_value = 0
        prescribed_rooms = {ca["prescribed_room"] for ca in self.course_assignments
                            if ca["prescribed_room"]}
        for room_id in prescribed_rooms:
            if room_id in rooms and rooms[room_id].total_priority > self.prescribed_room_value:
                self.prescribed_room_value = rooms[room_id].total_priority

        self.total_priority = self.raw_priority + self.top_student_value + self.prescribed_room_value

    @property
    def name(self):
        return f"{self.last_name}, {self.first_name}"

    def get_max_load(self, term):
        base = self.max_periods
        if term == "S1":
            if self.sixth_fy or self.sixth_s1:
                return base + 1
        elif term == "S2":
            if self.sixth_fy or self.sixth_s2:
                return base + 1
        return base


class Room:
    def __init__(self, room_id, capacity, available_periods, available_terms, shared):
        self.id = room_id
        self.capacity = capacity
        self.available_periods = available_periods  # list of period letters
        self.available_terms = available_terms  # list of terms or None (all)
        self.shared = shared

        self.raw_priority = 0
        self.total_priority = 0
        self.top_teacher_value = 0
        self.top_student_value = 0
        self.demand = 0
        self.lock_count = 0
        self.priority_history = {}

    def count_locks(self):
        locks = 0
        for p in PERIODS:
            if p not in self.available_periods:
                locks += 1
        if self.available_terms:
            for t in TERMS:
                if t not in self.available_terms:
                    locks += 1
        self.lock_count = locks

    def calculate_raw(self):
        self.count_locks()
        self.raw_priority = (self.lock_count * LOCK_WEIGHT) + (self.demand * ROOM_DEMAND_WEIGHT)

    def calculate_total(self, teachers, students):
        self.top_teacher_value = 0
        self.top_student_value = 0

        for t in teachers.values():
            prescribed_rooms = {ca["prescribed_room"] for ca in t.course_assignments
                                if ca["prescribed_room"]}
            if self.id in prescribed_rooms:
                if t.total_priority > self.top_teacher_value:
                    self.top_teacher_value = t.total_priority

        all_course_codes = set()
        for t in teachers.values():
            for ca in t.course_assignments:
                if ca["prescribed_room"] == self.id:
                    all_course_codes.add(ca["course_code"])

        for s in students.values():
            remaining = [cr for cr in s.course_requests if cr not in s.placed_courses]
            if any(cr in all_course_codes for cr in remaining):
                if s.total_priority > self.top_student_value:
                    self.top_student_value = s.total_priority

        self.total_priority = self.raw_priority + self.top_teacher_value + self.top_student_value


class Course:
    def __init__(self, code, title, department, credits, prescribed_term,
                 grade_levels, sections_needed, max_enrollment,
                 singleton, ap, grad_req, cohort, ncaa,
                 prerequisites, corequisites):
        self.code = code
        self.title = title
        self.department = department
        self.credits = credits
        self.prescribed_term = prescribed_term
        self.grade_levels = grade_levels
        self.sections_needed = sections_needed
        self.max_enrollment = max_enrollment
        self.singleton = singleton
        self.ap = ap
        self.grad_req = grad_req
        self.cohort = cohort
        self.ncaa = ncaa
        self.prerequisites = prerequisites
        self.corequisites = corequisites

    def calculate_raw(self):
        val = 0
        if self.ap:
            val += COURSE_CHAR_POINTS["ap"]
        if self.singleton:
            val += COURSE_CHAR_POINTS["singleton"]
        if self.grad_req:
            val += COURSE_CHAR_POINTS["grad_req"]
        if self.prescribed_term in ("S1", "S2"):
            val += COURSE_CHAR_POINTS["semester_only"]
        if self.cohort:
            val += COURSE_CHAR_POINTS["cohort_course"]
        if self.prescribed_term and self.prescribed_term != "FY":
            val += COURSE_CHAR_POINTS["prescribed_term"]
        return val


class Section:
    def __init__(self, course_code, section_num, teacher_id, prescribed_room,
                 prescribed_period, prescribed_term, prescribed_cohort):
        self.course_code = course_code
        self.section_num = section_num
        self.teacher_id = teacher_id
        self.prescribed_room = prescribed_room
        self.prescribed_period = prescribed_period
        self.prescribed_term = prescribed_term
        self.prescribed_cohort = prescribed_cohort

        self.assigned_room = prescribed_room
        self.assigned_period = None
        self.assigned_term = prescribed_term
        self.enrolled_students = []
        self.not_seated = []

        self.raw_priority = 0
        self.total_priority = 0
        self.top_student_value = 0
        self.teacher_total = 0
        self.room_total = 0
        self.placed = False
        self.co_schedule_group = None
        self.priority_history = {}

    @property
    def key(self):
        return f"{self.course_code}-{self.section_num}"

    def calculate_raw(self, course):
        val = course.calculate_raw()
        if self.co_schedule_group:
            val += COURSE_CHAR_POINTS["co_schedule"]
        self.raw_priority = val

    def calculate_total(self, students, teachers, rooms, course):
        self.top_student_value = 0
        for s in students.values():
            remaining = [cr for cr in s.course_requests if cr not in s.placed_courses]
            if self.course_code not in remaining:
                continue
            if s.total_priority > self.top_student_value:
                self.top_student_value = s.total_priority

        self.teacher_total = 0
        if self.teacher_id and self.teacher_id in teachers:
            self.teacher_total = teachers[self.teacher_id].total_priority

        self.room_total = 0
        room_id = self.assigned_room
        if room_id and room_id in rooms:
            self.room_total = rooms[room_id].total_priority

        self.total_priority = (self.raw_priority + self.top_student_value +
                               self.teacher_total + self.room_total)


class CoScheduleGroup:
    def __init__(self, name, course_codes, teacher_id, prescribed_room):
        self.name = name
        self.course_codes = course_codes
        self.teacher_id = teacher_id
        self.prescribed_room = prescribed_room
        self.section_keys = []
        self.raw_priority = 0
        self.total_priority = 0
        self.placed = False
        self.assigned_period = None
        self.assigned_term = None
        self.priority_history = {}

    def calculate_raw(self, sections, courses):
        val = 0
        for sk in self.section_keys:
            if sk in sections:
                val += sections[sk].raw_priority
        self.raw_priority = val

    def calculate_total(self, sections, students, teachers, rooms):
        top_student = 0
        for sk in self.section_keys:
            sec = sections.get(sk)
            if not sec:
                continue
            for s in students.values():
                remaining = [cr for cr in s.course_requests if cr not in s.placed_courses]
                if sec.course_code in remaining:
                    if s.total_priority > top_student:
                        top_student = s.total_priority

        teacher_total = 0
        if self.teacher_id and self.teacher_id in teachers:
            teacher_total = teachers[self.teacher_id].total_priority

        room_total = 0
        if self.prescribed_room and self.prescribed_room in rooms:
            room_total = rooms[self.prescribed_room].total_priority

        self.total_priority = self.raw_priority + top_student + teacher_total + room_total


# ===================================================================
# DATA LOADING
# ===================================================================

def yn(val):
    if val is None:
        return False
    s = str(val).strip().upper()
    if s in ("N/A", "NONE", ""):
        return False
    return s in ("Y", "YES", "TRUE", "1")


def parse_comma_list(val):
    if val is None:
        return []
    s = str(val).strip()
    if s.upper() in ("N/A", "NONE", ""):
        return []
    return [x.strip() for x in s.split(",") if x.strip()]


def load_students(path):
    """Load students from NEW template (11 columns on Sheet 1, 10 columns on Sheet 2).

    Sheet 1 — Student Profiles:
      A=Student ID, B=Last Name, C=First Name, D=Grade Level,
      E=NCAA, F=LEO II, G=LEO I, H=Academic Support, I=Pathway,
      J=Cohort Name, K=Cohort Locked

    Sheet 2 — Transcript History:
      A=Student ID, B=Academic Year, C=Course Code, D=Course Title,
      E=Department, F=Credits, G=Final Grade, H=Passed, I=Grade Level When Taken,
      J=Notes
    """
    wb = openpyxl.load_workbook(path)
    students = {}

    ws = wb["Student Profiles"]
    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        sid = vals[0]
        if sid is None:
            continue
        sid = str(sid).strip()
        if not sid or sid.upper() in ("STUDENT ID", "REQUIRED", "OPTIONAL"):
            continue

        s = Student(
            sid=sid,
            last_name=str(vals[1] or ""),
            first_name=str(vals[2] or ""),
            grade=int(vals[3]) if vals[3] else 9,
            ncaa=yn(vals[4]) if len(vals) > 4 else False,
            leo_ii=yn(vals[5]) if len(vals) > 5 else False,
            leo_i=yn(vals[6]) if len(vals) > 6 else False,
            academic_support=yn(vals[7]) if len(vals) > 7 else False,
            pathway=yn(vals[8]) if len(vals) > 8 else False,
            cohort_name=str(vals[9]).strip() if len(vals) > 9 and vals[9] else None,
            cohort_locked=yn(vals[10]) if len(vals) > 10 else False,
        )
        students[sid] = s

    if "Transcript History" in wb.sheetnames:
        ws2 = wb["Transcript History"]
        for row in ws2.iter_rows(min_row=2, values_only=False):
            vals = [c.value for c in row]
            sid = str(vals[0]).strip() if vals[0] else None
            if not sid or sid.upper() in ("STUDENT ID", "REQUIRED", "OPTIONAL"):
                continue
            if sid in students:
                students[sid].transcript.append({
                    "year": str(vals[1] or ""),
                    "course_code": str(vals[2] or ""),
                    "course_title": str(vals[3] or ""),
                    "department": str(vals[4] or ""),
                    "credits": vals[5],
                    "final_grade": vals[6],
                    "passed": yn(vals[7]),
                    "grade_level": vals[8],
                    "notes": str(vals[9] or ""),
                })

    return students


def load_course_requests(path, students):
    """Load course requests from NEW template (2 columns).

    A=Student ID, B=Course Code
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        sid = str(vals[0]).strip() if vals[0] else None
        if not sid or sid.upper() in ("STUDENT ID", "REQUIRED", "OPTIONAL"):
            continue
        course_code = str(vals[1]).strip() if len(vals) > 1 and vals[1] else None
        if sid and course_code and sid in students:
            if course_code not in students[sid].course_requests:
                students[sid].course_requests.append(course_code)


def load_teachers(path):
    """Load teachers from NEW template (17 columns on Sheet 1, 6 columns on Sheet 2).

    Sheet 1 — Teacher Profiles:
      A=Teacher ID, B=Last Name, C=First Name, D=Department,
      E=Max Teaching Periods, F=Approved 6th Period FY, G=Approved 6th Period S1 Only,
      H=Approved 6th Period S2 Only, I=Avail Period A, J=Avail Period B,
      K=Avail Period C, L=Avail Period D, M=Avail Period E, N=Avail Period F,
      O=Avail Period G, P=Special Student Population Teacher,
      Q=Approved for Co-Scheduled Sections

    Sheet 2 — Teacher-Course Assignments:
      A=Teacher ID, B=Course Code, C=Prescribed Room, D=Prescribed Period,
      E=Prescribed Term, F=Prescribed Cohort
    """
    wb = openpyxl.load_workbook(path)
    teachers = {}

    ws = wb["Teacher Profiles"] if "Teacher Profiles" in wb.sheetnames else wb[wb.sheetnames[0]]

    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        tid = vals[0]
        if tid is None:
            continue
        tid = str(tid).strip()
        if not tid or tid.upper() in ("TEACHER ID", "REQUIRED", "OPTIONAL"):
            continue

        avail = {}
        for i, p in enumerate(PERIODS):
            col_idx = 8 + i  # columns I through O (0-based 8-14)
            avail[p] = yn(vals[col_idx]) if col_idx < len(vals) else True

        t = Teacher(
            tid=tid,
            last_name=str(vals[1] or ""),
            first_name=str(vals[2] or ""),
            department=str(vals[3] or ""),
            max_periods=int(vals[4]) if len(vals) > 4 and vals[4] else 5,
            sixth_fy=yn(vals[5]) if len(vals) > 5 else False,
            sixth_s1=yn(vals[6]) if len(vals) > 6 else False,
            sixth_s2=yn(vals[7]) if len(vals) > 7 else False,
            avail=avail,
            ssp_teacher=str(vals[15] or "").strip() if len(vals) > 15 and vals[15] else "",
            co_schedule_approved=yn(vals[16]) if len(vals) > 16 else True,
        )
        teachers[tid] = t

    # Sheet 2: Teacher-Course Assignments
    sheet2_name = "Teacher-Course Assignments"
    if sheet2_name not in wb.sheetnames:
        for name in wb.sheetnames:
            if "assignment" in name.lower() or "course" in name.lower():
                sheet2_name = name
                break

    if sheet2_name in wb.sheetnames:
        ws2 = wb[sheet2_name]
        for row in ws2.iter_rows(min_row=2, values_only=False):
            vals = [c.value for c in row]
            tid = str(vals[0]).strip() if vals[0] else None
            if not tid or tid.upper() in ("TEACHER ID", "REQUIRED", "OPTIONAL"):
                continue
            course_code = str(vals[1]).strip() if vals[1] else None
            if not course_code:
                continue
            if tid in teachers:
                prescribed_room = str(vals[2]).strip() if len(vals) > 2 and vals[2] else None
                prescribed_period = str(vals[3]).strip().upper() if len(vals) > 3 and vals[3] else None
                prescribed_term = str(vals[4]).strip().upper() if len(vals) > 4 and vals[4] else None
                prescribed_cohort = str(vals[5]).strip() if len(vals) > 5 and vals[5] else None

                if prescribed_period and prescribed_period not in PERIODS:
                    prescribed_period = None
                if prescribed_term and prescribed_term not in TERMS:
                    prescribed_term = None

                teachers[tid].course_assignments.append({
                    "course_code": course_code,
                    "prescribed_room": prescribed_room,
                    "prescribed_period": prescribed_period,
                    "prescribed_term": prescribed_term,
                    "prescribed_cohort": prescribed_cohort,
                })

    return teachers


def load_rooms(path):
    """Load rooms from NEW template (5 columns).

    A=Room ID, B=Capacity, C=Available Periods, D=Available Terms, E=Shared Room
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rooms = {}

    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        rid = vals[0]
        if rid is None:
            continue
        rid = str(rid).strip()
        if not rid or rid.upper() in ("ROOM ID", "REQUIRED", "OPTIONAL"):
            continue

        avail_periods = parse_comma_list(vals[2]) if len(vals) > 2 and vals[2] else list(PERIODS)
        avail_terms = parse_comma_list(vals[3]) if len(vals) > 3 and vals[3] else None
        if avail_terms and not avail_terms:
            avail_terms = None

        r = Room(
            room_id=rid,
            capacity=int(vals[1]) if len(vals) > 1 and vals[1] else 30,
            available_periods=avail_periods,
            available_terms=avail_terms if avail_terms else None,
            shared=yn(vals[4]) if len(vals) > 4 else False,
        )
        rooms[rid] = r

    return rooms


def load_courses(path):
    """Load courses from NEW template (15 columns).

    A=Course Code, B=Course Title, C=Department, D=Credits,
    E=Prescribed Term, F=Grade Levels, G=Sections Needed,
    H=Max Enrollment per Section, I=Singleton, J=AP,
    K=Graduation Requirement, L=Cohort, M=NCAA,
    N=Prerequisites, O=Corequisites
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    courses = {}

    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        code = vals[0]
        if code is None:
            continue
        code = str(code).strip()
        if not code or code.upper() in ("REQUIRED", "OPTIONAL", "COURSE CODE"):
            continue

        term_raw = str(vals[4] or "FY").strip().upper() if len(vals) > 4 and vals[4] else "FY"
        if term_raw not in TERMS:
            term_raw = "FY"

        grad_req_raw = str(vals[10] or "").strip() if len(vals) > 10 and vals[10] else None
        if grad_req_raw and grad_req_raw.upper() in ("N/A", "NONE", ""):
            grad_req_raw = None

        cohort_raw = str(vals[11] or "").strip() if len(vals) > 11 and vals[11] else None
        if cohort_raw and cohort_raw.upper() in ("N/A", "NONE", ""):
            cohort_raw = None

        c = Course(
            code=code,
            title=str(vals[1] or ""),
            department=str(vals[2] or ""),
            credits=float(vals[3]) if len(vals) > 3 and vals[3] else 5,
            prescribed_term=term_raw,
            grade_levels=parse_comma_list(vals[5]) if len(vals) > 5 else [],
            sections_needed=int(vals[6]) if len(vals) > 6 and vals[6] else 1,
            max_enrollment=int(vals[7]) if len(vals) > 7 and vals[7] else 25,
            singleton=yn(vals[8]) if len(vals) > 8 else False,
            ap=yn(vals[9]) if len(vals) > 9 else False,
            grad_req=grad_req_raw,
            cohort=cohort_raw,
            ncaa=yn(vals[12]) if len(vals) > 12 else False,
            prerequisites=parse_comma_list(vals[13]) if len(vals) > 13 else [],
            corequisites=parse_comma_list(vals[14]) if len(vals) > 14 else [],
        )
        courses[code] = c

    return courses


def load_co_schedule_groups(path):
    """Load co-schedule groups from NEW template (4 columns, one row per course).

    A=Co-Schedule Group Name, B=Course Code, C=Teacher ID, D=Prescribed Room
    """
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    groups = {}

    for row in ws.iter_rows(min_row=2, values_only=False):
        vals = [c.value for c in row]
        name = vals[0]
        if name is None:
            continue
        name = str(name).strip()
        if not name or name.upper() in ("CO-SCHEDULE GROUP NAME", "REQUIRED", "OPTIONAL"):
            continue

        course_code = str(vals[1]).strip() if len(vals) > 1 and vals[1] else None
        teacher_id = str(vals[2]).strip() if len(vals) > 2 and vals[2] else None
        prescribed_room = str(vals[3]).strip() if len(vals) > 3 and vals[3] else None

        if not course_code:
            continue

        if name not in groups:
            groups[name] = {
                "name": name,
                "course_codes": [course_code],
                "teacher_id": teacher_id,
                "prescribed_room": prescribed_room,
            }
        else:
            groups[name]["course_codes"].append(course_code)
            if teacher_id and not groups[name]["teacher_id"]:
                groups[name]["teacher_id"] = teacher_id
            if prescribed_room and not groups[name]["prescribed_room"]:
                groups[name]["prescribed_room"] = prescribed_room

    return {
        name: CoScheduleGroup(
            name=g["name"],
            course_codes=g["course_codes"],
            teacher_id=g["teacher_id"],
            prescribed_room=g["prescribed_room"],
        )
        for name, g in groups.items()
    }


def load_semester_designations(path):
    if not os.path.exists(path):
        return {}
    with open(path) as f:
        data = json.load(f)
    result = {}
    for entry in data:
        code = str(entry.get("code", "")).strip()
        for sd in entry.get("section_details", []):
            sec_num = sd.get("sec", 1)
            sem = sd.get("semester", "")
            teacher = sd.get("teacher", "")
            result[(code, sec_num)] = {"semester": sem, "teacher": teacher}
    return result


# ===================================================================
# PRE-BUILD VALIDATION
# ===================================================================

def validate_pre_build(students, courses, teachers):
    errors = []
    warnings = []

    for sid, s in students.items():
        # Credit cap check
        total_credits = sum(
            courses[cr].credits for cr in s.course_requests if cr in courses
        )
        if total_credits > 35.0:
            errors.append(f"Student {sid} ({s.name}): {total_credits} credits requested, exceeds 35.0 cap")

        # Duplicate requests
        seen = set()
        for cr in s.course_requests:
            if cr in seen:
                warnings.append(f"Student {sid} ({s.name}): duplicate request for course {cr}")
            seen.add(cr)

        # Already completed
        completed = {t["course_code"] for t in s.transcript if t["passed"]}
        for cr in s.course_requests:
            if cr in completed:
                warnings.append(f"Student {sid} ({s.name}): requesting course {cr} already passed")

        # Prerequisites
        for cr in s.course_requests:
            if cr in courses:
                for prereq in courses[cr].prerequisites:
                    if prereq not in completed:
                        errors.append(
                            f"Student {sid} ({s.name}): course {cr} requires prerequisite {prereq} not completed"
                        )

        # NCAA validation
        if s.ncaa:
            for cr in s.course_requests:
                if cr in courses and not courses[cr].ncaa:
                    warnings.append(
                        f"Student {sid} ({s.name}): NCAA student requesting non-NCAA course {cr}"
                    )

        # Grade level eligibility
        for cr in s.course_requests:
            if cr in courses and courses[cr].grade_levels:
                if str(s.grade) not in courses[cr].grade_levels:
                    errors.append(
                        f"Student {sid} ({s.name}): grade {s.grade} not eligible for course {cr} "
                        f"(grades: {courses[cr].grade_levels})"
                    )

    # Teacher load validation
    for tid, t in teachers.items():
        s1_load = 0
        s2_load = 0
        for ca in t.course_assignments:
            term = ca.get("prescribed_term", "FY") or "FY"
            if term == "FY":
                s1_load += 1
                s2_load += 1
            elif term == "S1":
                s1_load += 1
            elif term == "S2":
                s2_load += 1
        max_s1 = t.get_max_load("S1")
        max_s2 = t.get_max_load("S2")
        if s1_load > max_s1:
            errors.append(f"Teacher {tid} ({t.name}): S1 load {s1_load} exceeds cap {max_s1}")
        if s2_load > max_s2:
            errors.append(f"Teacher {tid} ({t.name}): S2 load {s2_load} exceeds cap {max_s2}")

    return errors, warnings


# ===================================================================
# PRIORITY CALCULATIONS
# ===================================================================

def calculate_all_raw(students, teachers, rooms, courses, sections, co_groups):
    for s in students.values():
        s.calculate_raw()

    for t in teachers.values():
        t.calculate_raw()

    for r in rooms.values():
        r.demand = 0
    for sec in sections.values():
        if sec.assigned_room and sec.assigned_room in rooms:
            rooms[sec.assigned_room].demand += 1
    for r in rooms.values():
        r.calculate_raw()

    for sec in sections.values():
        if sec.course_code in courses:
            sec.calculate_raw(courses[sec.course_code])

    for cg in co_groups.values():
        cg.calculate_raw(sections, courses)


def calculate_course_request_priorities(students, courses, sections, teachers):
    for s in students.values():
        s.course_priority_values = {}
        for cr in s.course_requests:
            if cr in courses:
                course = courses[cr]
                val = 0
                categories = []
                if course.ap:
                    categories.append(COURSE_CHAR_POINTS["ap"])
                if course.singleton:
                    categories.append(COURSE_CHAR_POINTS["singleton"])
                if course.grad_req:
                    categories.append(COURSE_CHAR_POINTS["grad_req"])
                if course.prescribed_term in ("S1", "S2"):
                    categories.append(COURSE_CHAR_POINTS["semester_only"])
                if course.cohort:
                    categories.append(COURSE_CHAR_POINTS["cohort_course"])
                if course.prescribed_term and course.prescribed_term != "FY":
                    categories.append(COURSE_CHAR_POINTS["prescribed_term"])
                # Same course, multiple categories = highest value only
                val = max(categories) if categories else 0
                s.course_priority_values[cr] = val


def calculate_all_totals(students, teachers, rooms, courses, sections, co_groups):
    for s in students.values():
        s.calculate_total()

    for t in teachers.values():
        t.calculate_total(students, rooms)

    for r in rooms.values():
        r.calculate_total(teachers, students)

    for sec in sections.values():
        if sec.course_code in courses and not sec.placed:
            sec.calculate_total(students, teachers, rooms, courses[sec.course_code])

    for cg in co_groups.values():
        if not cg.placed:
            cg.calculate_total(sections, students, teachers, rooms)


# ===================================================================
# SECTION GENERATION
# ===================================================================

def generate_sections(courses, teachers, semester_designations):
    sections = {}
    teacher_course_map = defaultdict(list)
    for tid, t in teachers.items():
        for ca in t.course_assignments:
            cc = ca["course_code"]
            if cc:
                teacher_course_map[cc].append((tid, ca))

    for code, course in courses.items():
        assignments = teacher_course_map.get(code, [])

        for sec_num in range(1, course.sections_needed + 1):
            tid = None
            p_room = None
            p_period = None
            p_term = course.prescribed_term
            p_cohort = None

            sd_key = (code, sec_num)
            sd = semester_designations.get(sd_key, {})
            if sd.get("semester") and sd["semester"] not in ("Builder Choice", ""):
                sem = sd["semester"]
                if "S1" in sem or "Fall" in sem:
                    p_term = "S1"
                elif "S2" in sem or "Spring" in sem:
                    p_term = "S2"

            if assignments:
                if sec_num <= len(assignments):
                    tid, ca = assignments[sec_num - 1]
                    p_room = ca["prescribed_room"]
                    p_period = ca["prescribed_period"]
                    if ca["prescribed_term"]:
                        p_term = ca["prescribed_term"]
                    p_cohort = ca["prescribed_cohort"]
                else:
                    tid, ca = assignments[0]

            sec = Section(
                course_code=code,
                section_num=sec_num,
                teacher_id=tid,
                prescribed_room=p_room,
                prescribed_period=p_period,
                prescribed_term=p_term,
                prescribed_cohort=p_cohort,
            )
            sections[sec.key] = sec

    return sections


# ===================================================================
# SCHEDULE GRID
# ===================================================================

class ScheduleGrid:
    def __init__(self):
        self.grid = {}  # (period, term) -> list of section keys
        self.teacher_schedule = defaultdict(dict)  # tid -> {(period, term): section_key}
        self.room_schedule = defaultdict(dict)  # room_id -> {(period, term): section_key}
        self.student_schedule = defaultdict(dict)  # sid -> {(period, term): section_key}

    def is_slot_free_for_teacher(self, tid, period, term):
        if not tid:
            return True
        slots_to_check = [(period, term)]
        if term == "FY":
            slots_to_check = [(period, "FY"), (period, "S1"), (period, "S2")]
        elif term in ("S1", "S2"):
            slots_to_check = [(period, term), (period, "FY")]

        for slot in slots_to_check:
            if slot in self.teacher_schedule.get(tid, {}):
                return False
        return True

    def is_slot_free_for_room(self, room_id, period, term):
        if not room_id:
            return True
        slots_to_check = [(period, term)]
        if term == "FY":
            slots_to_check = [(period, "FY"), (period, "S1"), (period, "S2")]
        elif term in ("S1", "S2"):
            slots_to_check = [(period, term), (period, "FY")]

        for slot in slots_to_check:
            if slot in self.room_schedule.get(room_id, {}):
                return False
        return True

    def is_slot_free_for_student(self, sid, period, term):
        slots_to_check = [(period, term)]
        if term == "FY":
            slots_to_check = [(period, "FY"), (period, "S1"), (period, "S2")]
        elif term in ("S1", "S2"):
            slots_to_check = [(period, term), (period, "FY")]

        for slot in slots_to_check:
            if slot in self.student_schedule.get(sid, {}):
                return False
        return True

    def place_section(self, section_key, period, term, teacher_id, room_id):
        slot = (period, term)
        if slot not in self.grid:
            self.grid[slot] = []
        self.grid[slot].append(section_key)

        if teacher_id:
            self.teacher_schedule[teacher_id][slot] = section_key
            if term == "FY":
                self.teacher_schedule[teacher_id][(period, "S1")] = section_key
                self.teacher_schedule[teacher_id][(period, "S2")] = section_key

        if room_id:
            self.room_schedule[room_id][slot] = section_key
            if term == "FY":
                self.room_schedule[room_id][(period, "S1")] = section_key
                self.room_schedule[room_id][(period, "S2")] = section_key

    def seat_student(self, sid, period, term, section_key):
        slot = (period, term)
        self.student_schedule[sid][slot] = section_key
        if term == "FY":
            self.student_schedule[sid][(period, "S1")] = section_key
            self.student_schedule[sid][(period, "S2")] = section_key

    def get_valid_slots(self, teacher_id, room_id, term, teachers, rooms):
        valid = []
        periods_to_try = list(PERIODS)

        if teacher_id and teacher_id in teachers:
            t = teachers[teacher_id]
            periods_to_try = [p for p in periods_to_try if t.availability.get(p, True)]

        if room_id and room_id in rooms:
            r = rooms[room_id]
            periods_to_try = [p for p in periods_to_try if p in r.available_periods]

        for p in periods_to_try:
            if self.is_slot_free_for_teacher(teacher_id, p, term):
                if self.is_slot_free_for_room(room_id, p, term):
                    valid.append(p)

        return valid


# ===================================================================
# PLACEMENT ENGINE
# ===================================================================

class RunLog:
    def __init__(self):
        self.entries = []

    def record(self, run_number, section_key, period, term, section_total,
               top_student, teacher_total, room_total, seated, not_seated,
               raw_priority):
        self.entries.append({
            "run": run_number,
            "section": section_key,
            "period": period,
            "term": term,
            "course_section_raw": raw_priority,
            "course_section_total": section_total,
            "top_student_value": top_student,
            "teacher_total_value": teacher_total,
            "room_total_value": room_total,
            "students_seated": len(seated),
            "students_not_seated": len(not_seated),
        })


class ClashReport:
    def __init__(self):
        self.clashes = []

    def add(self, section_key, reason, blocking_section=None, options=None):
        self.clashes.append({
            "section": section_key,
            "reason": reason,
            "blocking_section": blocking_section,
            "resolution_options": options or [],
        })


def find_room_for_section(section, period, term, rooms, grid):
    if section.assigned_room:
        if grid.is_slot_free_for_room(section.assigned_room, period, term):
            return section.assigned_room
        return None

    best_room = None
    best_capacity = float("inf")

    for rid, room in rooms.items():
        if period not in room.available_periods:
            continue
        if room.available_terms and term not in room.available_terms:
            continue
        if not grid.is_slot_free_for_room(rid, period, term):
            continue
        if room.capacity < best_capacity:
            best_room = rid
            best_capacity = room.capacity

    return best_room


def seat_students(section, students, courses, grid, period, term):
    course = courses.get(section.course_code)
    if not course:
        return [], []

    eligible = []
    for sid, s in students.items():
        remaining = [cr for cr in s.course_requests if cr not in s.placed_courses]
        if section.course_code not in remaining:
            continue
        if not grid.is_slot_free_for_student(sid, period, term):
            continue
        eligible.append((s.total_priority, sid))

    eligible.sort(key=lambda x: -x[0])

    cap = course.max_enrollment
    seated = []
    not_seated = []

    if section.prescribed_cohort:
        for prio, sid in eligible:
            s = students[sid]
            if s.cohort_name == section.prescribed_cohort:
                seated.append(sid)
                grid.seat_student(sid, period, term, section.key)
                s.placed_courses.add(section.course_code)

        for prio, sid in eligible:
            if sid not in seated and len(seated) < cap:
                s = students[sid]
                if s.cohort_name and s.cohort_name != section.prescribed_cohort:
                    continue
                seated.append(sid)
                grid.seat_student(sid, period, term, section.key)
                s.placed_courses.add(section.course_code)
            elif sid not in seated:
                not_seated.append(sid)
    else:
        for prio, sid in eligible:
            if len(seated) < cap:
                seated.append(sid)
                grid.seat_student(sid, period, term, section.key)
                students[sid].placed_courses.add(section.course_code)
            else:
                not_seated.append(sid)

    return seated, not_seated


def run_placement(students, teachers, rooms, courses, sections, co_groups,
                  grid, run_log, clash_report):
    run_number = 0

    co_group_sections = set()
    for cg in co_groups.values():
        for sk in cg.section_keys:
            co_group_sections.add(sk)

    while True:
        unplaced_sections = {k: s for k, s in sections.items()
                             if not s.placed and k not in co_group_sections}
        unplaced_groups = {k: g for k, g in co_groups.items() if not g.placed}

        if not unplaced_sections and not unplaced_groups:
            break

        # Build combined ranking: sections and co-schedule groups
        candidates = []
        for k, sec in unplaced_sections.items():
            candidates.append(("section", k, sec.total_priority, sec.raw_priority))
        for k, cg in unplaced_groups.items():
            candidates.append(("group", k, cg.total_priority, cg.raw_priority))

        # Sort: highest total first, tiebreaker = highest raw
        candidates.sort(key=lambda x: (-x[2], -x[3]))

        if not candidates:
            break

        kind, key, total, raw = candidates[0]
        run_number += 1

        if kind == "section":
            sec = sections[key]
            placed = place_single_section(
                sec, students, teachers, rooms, courses, grid,
                run_log, clash_report, run_number
            )
            if placed:
                recalculate_after_placement(
                    students, teachers, rooms, courses, sections, co_groups
                )

        elif kind == "group":
            cg = co_groups[key]
            placed = place_co_schedule_group(
                cg, sections, students, teachers, rooms, courses,
                grid, run_log, clash_report, run_number
            )
            if placed:
                recalculate_after_placement(
                    students, teachers, rooms, courses, sections, co_groups
                )

        # Store priority history for this run
        for s in students.values():
            s.priority_history[run_number] = s.total_priority
        for t in teachers.values():
            t.priority_history[run_number] = t.total_priority
        for r in rooms.values():
            r.priority_history[run_number] = r.total_priority

        if run_number > 5000:
            break


def place_single_section(sec, students, teachers, rooms, courses, grid,
                         run_log, clash_report, run_number):
    term = sec.assigned_term or "FY"
    room_id = sec.assigned_room

    if sec.prescribed_period:
        period = sec.prescribed_period
        if not grid.is_slot_free_for_teacher(sec.teacher_id, period, term):
            clash_report.add(
                sec.key,
                f"Prescribed period {period} conflicts with teacher {sec.teacher_id} schedule",
                options=[
                    {"option": "A", "change": f"Move to a different period",
                     "side_effects": "Teacher must be available"},
                ]
            )
            sec.placed = True
            return False

        if room_id and not grid.is_slot_free_for_room(room_id, period, term):
            new_room = find_room_for_section(
                Section(sec.course_code, sec.section_num, sec.teacher_id,
                        None, sec.prescribed_period, sec.prescribed_term,
                        sec.prescribed_cohort),
                period, term, rooms, grid
            )
            if new_room:
                room_id = new_room
            else:
                clash_report.add(
                    sec.key,
                    f"Prescribed room {sec.assigned_room} not available at period {period} term {term}",
                )
                sec.placed = True
                return False

        if not room_id:
            room_id = find_room_for_section(sec, period, term, rooms, grid)
        if not room_id:
            clash_report.add(sec.key, f"No room available at period {period} term {term}")
            sec.placed = True
            return False

        grid.place_section(sec.key, period, term, sec.teacher_id, room_id)
        sec.assigned_period = period
        sec.assigned_room = room_id
        sec.assigned_term = term
        sec.placed = True

        seated, not_seated = seat_students(sec, students, courses, grid, period, term)
        sec.enrolled_students = seated
        sec.not_seated = not_seated

        run_log.record(
            run_number, sec.key, period, term, sec.total_priority,
            sec.top_student_value, sec.teacher_total, sec.room_total,
            seated, not_seated, sec.raw_priority
        )
        return True

    # No prescribed period — find best slot
    valid_periods = grid.get_valid_slots(sec.teacher_id, room_id, term, teachers, rooms)

    if not valid_periods and not room_id:
        for p in PERIODS:
            if grid.is_slot_free_for_teacher(sec.teacher_id, p, term):
                r = find_room_for_section(sec, p, term, rooms, grid)
                if r:
                    valid_periods.append(p)

    if not valid_periods:
        clash_report.add(
            sec.key,
            f"No valid period/room combination available for term {term}",
            options=[
                {"option": "A", "change": "Check teacher availability and room constraints"},
            ]
        )
        sec.placed = True
        return False

    best_period = valid_periods[0]

    if not room_id:
        room_id = find_room_for_section(sec, best_period, term, rooms, grid)
    if not room_id:
        for p in valid_periods[1:]:
            room_id = find_room_for_section(sec, p, term, rooms, grid)
            if room_id:
                best_period = p
                break

    if not room_id:
        clash_report.add(sec.key, f"No room available for any valid period in term {term}")
        sec.placed = True
        return False

    grid.place_section(sec.key, best_period, term, sec.teacher_id, room_id)
    sec.assigned_period = best_period
    sec.assigned_room = room_id
    sec.assigned_term = term
    sec.placed = True

    seated, not_seated = seat_students(sec, students, courses, grid, best_period, term)
    sec.enrolled_students = seated
    sec.not_seated = not_seated

    run_log.record(
        run_number, sec.key, best_period, term, sec.total_priority,
        sec.top_student_value, sec.teacher_total, sec.room_total,
        seated, not_seated, sec.raw_priority
    )
    return True


def place_co_schedule_group(cg, sections, students, teachers, rooms, courses,
                            grid, run_log, clash_report, run_number):
    term_candidates = list(TERMS)
    room_id = cg.prescribed_room
    teacher_id = cg.teacher_id

    best_placement = None
    for term in term_candidates:
        valid_periods = grid.get_valid_slots(teacher_id, room_id, term, teachers, rooms)
        if not room_id:
            expanded = []
            for p in PERIODS:
                if grid.is_slot_free_for_teacher(teacher_id, p, term):
                    r = find_room_for_section(
                        Section("", 0, teacher_id, None, None, term, None),
                        p, term, rooms, grid
                    )
                    if r:
                        expanded.append((p, r))
            if expanded:
                best_placement = (expanded[0][0], term, expanded[0][1])
                break
        elif valid_periods:
            best_placement = (valid_periods[0], term, room_id)
            break

    if not best_placement:
        clash_report.add(
            f"CoSchedule:{cg.name}",
            f"No valid period/term/room for co-schedule group",
        )
        cg.placed = True
        return False

    period, term, final_room = best_placement
    cg.assigned_period = period
    cg.assigned_term = term
    cg.placed = True

    for sk in cg.section_keys:
        sec = sections.get(sk)
        if not sec:
            continue
        grid.place_section(sk, period, term, teacher_id, final_room)
        sec.assigned_period = period
        sec.assigned_room = final_room
        sec.assigned_term = term
        sec.placed = True

        seated, not_seated = seat_students(sec, students, courses, grid, period, term)
        sec.enrolled_students = seated
        sec.not_seated = not_seated

        run_log.record(
            run_number, sk, period, term, sec.total_priority,
            sec.top_student_value, sec.teacher_total, sec.room_total,
            seated, not_seated, sec.raw_priority
        )

    return True


def recalculate_after_placement(students, teachers, rooms, courses, sections, co_groups):
    calculate_course_request_priorities(students, courses, sections, teachers)
    calculate_all_totals(students, teachers, rooms, courses, sections, co_groups)


# ===================================================================
# OUTPUT
# ===================================================================

def generate_output(sections, students, courses, teachers, rooms,
                    run_log, clash_report, output_dir):
    os.makedirs(output_dir, exist_ok=True)

    # Master schedule
    schedule = []
    for key, sec in sorted(sections.items()):
        course = courses.get(sec.course_code, None)
        teacher = teachers.get(sec.teacher_id, None) if sec.teacher_id else None
        schedule.append({
            "section_key": key,
            "course_code": sec.course_code,
            "course_title": course.title if course else "",
            "department": course.department if course else "",
            "section_num": sec.section_num,
            "teacher_id": sec.teacher_id,
            "teacher_name": teacher.name if teacher else "",
            "room": sec.assigned_room,
            "period": sec.assigned_period,
            "term": sec.assigned_term,
            "enrolled": len(sec.enrolled_students),
            "cap": course.max_enrollment if course else 0,
            "raw_priority": sec.raw_priority,
            "total_priority": sec.total_priority,
            "placed": sec.placed,
        })

    with open(os.path.join(output_dir, "master_schedule.json"), "w") as f:
        json.dump(schedule, f, indent=2)

    # Run log
    with open(os.path.join(output_dir, "run_log.json"), "w") as f:
        json.dump(run_log.entries, f, indent=2)

    # Clash report
    with open(os.path.join(output_dir, "clash_report.json"), "w") as f:
        json.dump(clash_report.clashes, f, indent=2)

    # Student placement summary
    student_summary = []
    for sid, s in sorted(students.items()):
        requested = len(s.course_requests)
        placed = len(s.placed_courses)
        student_summary.append({
            "student_id": sid,
            "name": s.name,
            "grade": s.grade,
            "raw_priority": s.raw_priority,
            "courses_requested": requested,
            "courses_placed": placed,
            "placement_rate": round(placed / requested * 100, 1) if requested else 0,
            "unplaced": [cr for cr in s.course_requests if cr not in s.placed_courses],
        })

    with open(os.path.join(output_dir, "student_summary.json"), "w") as f:
        json.dump(student_summary, f, indent=2)

    # Priority snapshots
    priority_data = {
        "students": {sid: {"raw": s.raw_priority, "final_total": s.total_priority,
                           "history": s.priority_history}
                     for sid, s in students.items()},
        "teachers": {tid: {"raw": t.raw_priority, "final_total": t.total_priority,
                           "lock_count": t.lock_count, "history": t.priority_history}
                     for tid, t in teachers.items()},
        "rooms": {rid: {"raw": r.raw_priority, "final_total": r.total_priority,
                        "demand": r.demand, "history": r.priority_history}
                  for rid, r in rooms.items()},
    }

    with open(os.path.join(output_dir, "priority_snapshots.json"), "w") as f:
        json.dump(priority_data, f, indent=2)

    return schedule


# ===================================================================
# MAIN
# ===================================================================

def main():
    template_dir = "templates"
    output_dir = "engine_output"

    print("=" * 60)
    print("MASTER SCHEDULE BUILDER — DON BOSCO PREP 2026-27")
    print("Engine v4: Built from DATA_STRUCTURE.md")
    print("=" * 60)

    # --- LOAD DATA ---
    print("\n[1/9] Loading students...")
    students = load_students(os.path.join(template_dir, "Template_8_Student_Profiles.xlsx"))
    print(f"       Loaded {len(students)} students")

    print("[2/9] Loading course requests...")
    load_course_requests(
        os.path.join(template_dir, "Template_2_Student_Course_Requests.xlsx"),
        students
    )
    total_requests = sum(len(s.course_requests) for s in students.values())
    print(f"       Loaded {total_requests} course requests")

    print("[3/9] Loading teachers...")
    teachers = load_teachers(
        os.path.join(template_dir, "Template_6_Teacher_Profiles.xlsx"),
    )
    print(f"       Loaded {len(teachers)} teachers")

    print("[4/9] Loading rooms...")
    rooms = load_rooms(os.path.join(template_dir, "Template_9_Room_Profiles.xlsx"))
    print(f"       Loaded {len(rooms)} rooms")

    print("[5/9] Loading courses...")
    courses = load_courses(os.path.join(template_dir, "Template_7_Course_Profiles.xlsx"))
    print(f"       Loaded {len(courses)} courses")

    print("[6/9] Loading co-schedule groups...")
    co_groups = load_co_schedule_groups(
        os.path.join(template_dir, "Template_4_CoSchedule_Groups.xlsx")
    )
    print(f"       Loaded {len(co_groups)} co-schedule groups")

    print("[7/9] Loading semester designations...")
    sem_path = "semester_designations.json"
    semester_designations = load_semester_designations(sem_path)
    print(f"       Loaded {len(semester_designations)} section designations")

    # --- PRE-BUILD VALIDATION ---
    print("\n[8/9] Running pre-build validation...")
    errors, warnings = validate_pre_build(students, courses, teachers)
    if warnings:
        print(f"       {len(warnings)} warnings found")
        for w in warnings[:10]:
            print(f"       WARNING: {w}")
        if len(warnings) > 10:
            print(f"       ... and {len(warnings) - 10} more")
    if errors:
        print(f"       {len(errors)} ERRORS found")
        for e in errors[:10]:
            print(f"       ERROR: {e}")
        if len(errors) > 10:
            print(f"       ... and {len(errors) - 10} more")

    # --- GENERATE SECTIONS ---
    print("\n[9/9] Generating sections...")
    sections = generate_sections(courses, teachers, semester_designations)
    print(f"       Generated {len(sections)} sections")

    # Link co-schedule groups to sections
    for cg in co_groups.values():
        for cc in cg.course_codes:
            for sk, sec in sections.items():
                if sec.course_code == cc:
                    cg.section_keys.append(sk)
                    sec.co_schedule_group = cg.name

    # --- CALCULATE RAW PRIORITIES ---
    print("\n--- CALCULATING RAW PRIORITIES ---")
    calculate_all_raw(students, teachers, rooms, courses, sections, co_groups)

    student_raws = sorted([(s.raw_priority, s.id, s.name) for s in students.values()], reverse=True)
    print(f"Student Raw range: {student_raws[-1][0]} to {student_raws[0][0]}")
    print(f"  Highest: {student_raws[0][2]} (ID: {student_raws[0][1]}) = {student_raws[0][0]}")

    teacher_raws = sorted([(t.raw_priority, t.id, t.name) for t in teachers.values()], reverse=True)
    print(f"Teacher Raw range: {teacher_raws[-1][0]} to {teacher_raws[0][0]}")
    if teacher_raws:
        print(f"  Highest: {teacher_raws[0][2]} (ID: {teacher_raws[0][1]}) = {teacher_raws[0][0]}, locks = {teachers[teacher_raws[0][1]].lock_count}")

    section_raws = sorted([(s.raw_priority, s.key) for s in sections.values()], reverse=True)
    print(f"Section Raw range: {section_raws[-1][0]} to {section_raws[0][0]}")
    print(f"  Highest: {section_raws[0][1]} = {section_raws[0][0]}")

    # --- CALCULATE TOTAL PRIORITIES ---
    print("\n--- CALCULATING TOTAL PRIORITIES ---")
    calculate_course_request_priorities(students, courses, sections, teachers)
    calculate_all_totals(students, teachers, rooms, courses, sections, co_groups)

    section_totals = sorted([(s.total_priority, s.key) for s in sections.values()], reverse=True)
    print(f"Section Total range: {section_totals[-1][0]} to {section_totals[0][0]}")
    print(f"  Top 5 sections to place first:")
    for total, key in section_totals[:5]:
        sec = sections[key]
        course = courses.get(sec.course_code)
        title = course.title if course else "?"
        print(f"    {key} ({title}) — Total: {total}, Raw: {sec.raw_priority}")

    # --- RUN PLACEMENT ---
    print("\n--- RUNNING PLACEMENT ENGINE ---")
    grid = ScheduleGrid()
    rlog = RunLog()
    clashes = ClashReport()

    run_placement(students, teachers, rooms, courses, sections, co_groups,
                  grid, rlog, clashes)

    placed_count = sum(1 for s in sections.values() if s.placed and s.assigned_period)
    total_sections = len(sections)
    print(f"\nPlacement complete: {placed_count}/{total_sections} sections placed")

    if clashes.clashes:
        print(f"Clashes found: {len(clashes.clashes)}")

    total_seated = sum(len(s.enrolled_students) for s in sections.values())
    total_not_seated = sum(len(s.not_seated) for s in sections.values())
    print(f"Students seated: {total_seated}")
    print(f"Students not seated: {total_not_seated}")

    # --- GENERATE OUTPUT ---
    print(f"\n--- GENERATING OUTPUT TO {output_dir}/ ---")
    generate_output(sections, students, courses, teachers, rooms, rlog, clashes, output_dir)
    print("Output files:")
    print("  master_schedule.json")
    print("  run_log.json")
    print("  clash_report.json")
    print("  student_summary.json")
    print("  priority_snapshots.json")

    print("\n" + "=" * 60)
    print("ENGINE BUILD COMPLETE")
    print("=" * 60)


if __name__ == "__main__":
    main()
